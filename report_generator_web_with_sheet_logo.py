
import streamlit as st
import pandas as pd
from docx import Document
import io
import os
import tempfile

st.set_page_config(page_title="ระบบรายงานผลการเรียน", layout="wide")

# โลโก้โรงเรียน (ปรับเปลี่ยนไฟล์เป็นโลโก้จริงได้)
st.image("https://upload.wikimedia.org/wikipedia/commons/thumb/4/45/Logo_School_example.svg/1200px-Logo_School_example.svg.png", width=100)
st.title("โรงเรียนแสงทองวิทยา")
st.markdown("### 📑 ระบบสร้างรายงานผลการเรียน รายบุคคล / หลายคน")

# อัปโหลดไฟล์ Excel
excel_file = st.file_uploader("📥 เลือกไฟล์ Excel (.xlsx)", type=["xlsx"])
template_file = st.file_uploader("📄 เลือกไฟล์ Word Template (.docx)", type=["docx"])

sheet_name = None
student_id_column = None
df = None

if excel_file:
    xls = pd.ExcelFile(excel_file)
    sheet_name = st.selectbox("เลือก Sheet ที่มีข้อมูล", xls.sheet_names)
    df = xls.parse(sheet_name)
    st.subheader("🧾 Preview หัวตาราง")
    st.dataframe(df.head())
    student_id_column = st.selectbox("เลือกคอลัมน์รหัสประจำตัวนักเรียน", df.columns.tolist())

report_type = st.radio("เลือกรายงาน", ["ทั้งหมด", "เฉพาะนักเรียนที่เลือก"])
output_format = st.selectbox("เลือกรูปแบบไฟล์", ["docx", "pdf"])
student_ids = st.text_input("กรอกเลขประจำตัว (คั่น comma)")

st.subheader("📝 รายละเอียดหัวรายงาน")
col1, col2 = st.columns(2)
with col1:
    term = st.text_input("ภาคเรียนที่", value="2")
    grade = st.text_input("ระดับชั้น", value="มัธยมศึกษาปีที่ 1/9")
with col2:
    year = st.text_input("ปีการศึกษา", value="2566")
    program = st.text_input("โปรแกรม", value="SME แสงทอง")

def replace_placeholders(doc, replacements):
    for p in doc.paragraphs:
        for key, val in replacements.items():
            if key in p.text:
                inline = p.runs
                for i in range(len(inline)):
                    if key in inline[i].text:
                        inline[i].text = inline[i].text.replace(key, val)
    return doc

if st.button("🚀 สร้างรายงาน"):
    if df is not None and template_file and student_id_column:
        df[student_id_column] = df[student_id_column].astype(str)

        if report_type == "ทั้งหมด":
            selected_students = df
        else:
            ids = [x.strip() for x in student_ids.split(",") if x.strip()]
            selected_students = df[df[student_id_column].isin(ids)]

        st.subheader("📊 รายชื่อนักเรียนที่กำลังสร้างรายงาน")
        st.dataframe(selected_students[[student_id_column, 'ชื่อ', 'นามสกุล']])

        for _, student in selected_students.iterrows():
            doc = Document(template_file)

            header_replacements = {
                "ภาคเรียนที่ 2": f"ภาคเรียนที่ {term}",
                "ปีการศึกษา 2566": f"ปีการศึกษา {year}",
                "มัธยมศึกษาปีที่ 1/9": grade,
                "SME แสงทอง": program
            }
            doc = replace_placeholders(doc, header_replacements)

            replace_dict = {
                '«title»': str(student.get('คำนำหน้า', '')),
                '«name»': str(student.get('ชื่อ', '')),
                '«last»': str(student.get('นามสกุล', '')),
                '«id»': str(student.get(student_id_column, '')),
                '«gt1»': str(student.get('ท21102', '')),
                '«gt2»': str(student.get('ค21102', '')),
                '«gt3»': str(student.get('ว21102', '')),
                '«gt4»': str(student.get('ส21103', '')),
                '«gt5»': str(student.get('ส21104', '')),
                '«gt6»': str(student.get('พ21102', '')),
                '«gt7»': str(student.get('ศ21102', '')),
                '«gt8»': str(student.get('ง21102', '')),
                '«gt9»': str(student.get('อ21102', '')),
                '«pt1»': str(student.get('ว21282', '')),
                '«pt2»': str(student.get('ค21202', '')),
                '«pt3»': str(student.get('ว21204', '')),
                '«pt4»': str(student.get('อ21208', '')),
                '«pt5»': str(student.get('อ21210', '')),
                '«pt6»': str(student.get('อ21212', '')),
                '«pt7»': str(student.get('ส21202', '')),
                '«grade2»': str(student.get('GPA', ''))
            }

            doc = replace_placeholders(doc, replace_dict)

            name = f"{student[student_id_column]}_{student.get('ชื่อ', 'student')}".replace(" ", "")
            buffer = io.BytesIO()

            if output_format == "docx":
                doc.save(buffer)
                buffer.seek(0)
                st.download_button(
                    label=f"📄 ดาวน์โหลด: รายงาน_{name}.docx",
                    data=buffer,
                    file_name=f"รายงาน_{name}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )
            elif output_format == "pdf":
                with tempfile.TemporaryDirectory() as tmpdir:
                    docx_path = os.path.join(tmpdir, f"{name}.docx")
                    pdf_path = os.path.join(tmpdir, f"{name}.pdf")
                    doc.save(docx_path)
                    try:
                        from docx2pdf import convert
                        convert(docx_path, pdf_path)
                        with open(pdf_path, "rb") as pdf_file:
                            st.download_button(
                                label=f"📄 ดาวน์โหลด: รายงาน_{name}.pdf",
                                data=pdf_file,
                                file_name=f"รายงาน_{name}.pdf",
                                mime="application/pdf"
                            )
                    except Exception:
                        st.error("❌ ไม่สามารถแปลงเป็น PDF ได้ (ระบบอาจไม่รองรับ docx2pdf)")
    else:
        st.error("⚠️ กรุณาอัปโหลดไฟล์ Excel, Word Template และเลือก Sheet/คอลัมน์ให้ครบ")
